import tkinter as tk
from tkinter import ttk, simpledialog
import openpyxl
from openpyxl import load_workbook
from datetime import datetime
import os

arquivo_excel = "manutencao.xlsx"
equipamentos = ["RJ45", "Cabo Power", "AC Adapter", "HDMI", "VGA", "Multiviewer", "KVM"]
lugares = ["RUNIN", "AVT"]

if not os.path.exists(arquivo_excel):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Data/Hora", "Equipamento", "Quantidade", "Motivo", "Lugar", "Linha", "Trave", "Ponto"])
    wb.save(arquivo_excel)

root = tk.Tk()
root.title("Registro de Manutenção")
root.grid_columnconfigure(1, weight=1)

equipamento_var = tk.StringVar()
lugar_var = tk.StringVar()

FIELD_WIDTH = 26

quantidade_entry = tk.Entry(root, width=FIELD_WIDTH)
motivo_entry = tk.Entry(root, width=FIELD_WIDTH)
linha_entry = tk.Entry(root, width=FIELD_WIDTH)
trave_entry = tk.Entry(root, width=FIELD_WIDTH)
ponto_entry = tk.Entry(root, width=FIELD_WIDTH)


def set_status(msg, focus_widget=None):
    status_label.config(text=msg)
    if focus_widget is not None:
        focus_widget.focus_set()


def limpar_campos():
    equipamento_var.set("")
    lugar_var.set("")
    quantidade_entry.delete(0, tk.END)
    motivo_entry.delete(0, tk.END)
    linha_entry.delete(0, tk.END)
    trave_entry.delete(0, tk.END)
    ponto_entry.delete(0, tk.END)
    equipamento_menu.focus_set()


def parse_quantidade():
    q_raw = quantidade_entry.get().strip()
    if q_raw == "":
        return 1, False  # (valor, foi_preenchida?)
    if not q_raw.isdigit():
        raise ValueError("Quantidade deve ser um número inteiro!")
    q = int(q_raw)
    if q < 1:
        raise ValueError("Quantidade deve ser >= 1!")
    return q, True


def parse_int_optional(entry_widget, nome_campo):
    raw = entry_widget.get().strip()
    if raw == "":
        return ""
    try:
        return int(raw)
    except ValueError:
        raise ValueError(f"{nome_campo} deve ser um número inteiro!")


def parse_pontos():
    pontos_raw = ponto_entry.get().strip()
    if not pontos_raw:
        return []

    pontos_list = []
    for p in pontos_raw.split(","):
        p = p.strip()
        if not p:
            continue
        if p.isdigit():
            pontos_list.append(int(p))
        else:
            raise ValueError(f"Ponto inválido: {p}")

    return pontos_list

def inserir_dados():
    equipamento = equipamento_var.get().strip()
    if not equipamento:
        set_status("Escolha ou digite um equipamento!", equipamento_menu)
        return

    motivo = motivo_entry.get().strip()
    lugar = lugar_var.get().strip()
    try:
        quantidade, quantidade_preenchida = parse_quantidade()
    except ValueError as e:
        set_status(str(e), quantidade_entry)
        return
    
    try:
        linha = parse_int_optional(linha_entry, "Linha")
        trave = parse_int_optional(trave_entry, "Trave")
    except ValueError as e:
        msg = str(e)
        if "Linha" in msg:
            set_status(msg, linha_entry)
        else:
            set_status(msg, trave_entry)
        return
    
    try:
        pontos_list = parse_pontos()
    except ValueError as e:
        set_status(str(e), ponto_entry)
        return
    
    if pontos_list and quantidade_preenchida:
        if quantidade != len(pontos_list):
            set_status(
                f"Erro: Quantidade ({quantidade}) precisa ser igual ao nº de pontos ({len(pontos_list)}).",
                quantidade_entry
            )
            return

    agora = datetime.now().strftime("%d/%m/%Y %H:%M:%S")

    wb = load_workbook(arquivo_excel)
    ws = wb.active

    if pontos_list:
        for ponto in pontos_list:
            ws.append([agora, equipamento, 1, motivo, lugar, linha, trave, ponto])
        wb.save(arquivo_excel)
        set_status(f"{len(pontos_list)} registros adicionados (quantidade 1 por ponto).")
    else:
        ws.append([agora, equipamento, quantidade, motivo, lugar, linha, trave, ""])
        wb.save(arquivo_excel)
        set_status(f"1 registro adicionado (quantidade {quantidade}).")

    limpar_campos()

def adicionar_equipamento():
    novo = simpledialog.askstring("Novo Equipamento", "Digite o nome do equipamento:")
    if not novo:
        return
    novo = novo.strip()
    if not novo:
        return

    if novo in equipamentos:
        set_status(f"{novo} já existe na lista.")
        return

    equipamentos.append(novo)
    equipamento_menu["values"] = equipamentos
    equipamento_var.set(novo)
    set_status(f"{novo} adicionado à lista de equipamentos!")

PADX = 8
PADY = 4

ttk.Label(root, text="Equipamento").grid(row=0, column=0, padx=PADX, pady=PADY, sticky="e")

equipamento_menu = ttk.Combobox(
    root, textvariable=equipamento_var, values=equipamentos, width=FIELD_WIDTH, state="normal"
)
equipamento_menu.grid(row=0, column=1, padx=PADX, pady=PADY, sticky="ew")

btn_adicionar_equip = ttk.Button(root, text="Adicionar Equipamento", command=adicionar_equipamento)
btn_adicionar_equip.grid(row=0, column=2, padx=PADX, pady=PADY, sticky="w")

ttk.Label(root, text="Quantidade").grid(row=1, column=0, padx=PADX, pady=PADY, sticky="e")
quantidade_entry.grid(row=1, column=1, padx=PADX, pady=PADY, sticky="ew")

ttk.Label(root, text="Motivo").grid(row=2, column=0, padx=PADX, pady=PADY, sticky="e")
motivo_entry.grid(row=2, column=1, padx=PADX, pady=PADY, sticky="ew")

ttk.Label(root, text="Lugar").grid(row=3, column=0, padx=PADX, pady=PADY, sticky="e")

lugar_menu = ttk.Combobox(
    root, textvariable=lugar_var, values=lugares, width=FIELD_WIDTH, state="normal"
)
lugar_menu.grid(row=3, column=1, padx=PADX, pady=PADY, sticky="ew")

ttk.Label(root, text="Linha").grid(row=4, column=0, padx=PADX, pady=PADY, sticky="e")
linha_entry.grid(row=4, column=1, padx=PADX, pady=PADY, sticky="ew")

ttk.Label(root, text="Trave").grid(row=5, column=0, padx=PADX, pady=PADY, sticky="e")
trave_entry.grid(row=5, column=1, padx=PADX, pady=PADY, sticky="ew")

ttk.Label(root, text="Ponto (Virgula para linhas adicionais)").grid(row=6, column=0, padx=PADX, pady=PADY, sticky="e")
ponto_entry.grid(row=6, column=1, padx=PADX, pady=PADY, sticky="ew")

btn_inserir = ttk.Button(root, text="Inserir", command=inserir_dados)
btn_inserir.grid(row=7, column=0, columnspan=3, padx=PADX, pady=10)

status_label = tk.Label(root, text="")
status_label.grid(row=8, column=0, columnspan=3, padx=PADX, pady=PADY, sticky="ew")

tab_order = [
    equipamento_menu,
    btn_adicionar_equip,
    quantidade_entry,
    motivo_entry,
    lugar_menu,
    linha_entry,
    trave_entry,
    ponto_entry,
    btn_inserir
]


def focus_next(event):
    widget = event.widget
    try:
        idx = tab_order.index(widget)
        next_widget = tab_order[(idx + 1) % len(tab_order)]
        next_widget.focus_set()
    except ValueError:
        pass
    return "break"


for w in tab_order:
    w.bind("<Tab>", focus_next)

btn_inserir.bind("<Return>", lambda e: inserir_dados())

equipamento_menu.focus_set()
root.mainloop()